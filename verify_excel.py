import openpyxl

def verify_wb(path):
    wb = openpyxl.load_workbook(path)
    print(f"=== Verifying {path} ===")
    print("Sheets in workbook:", wb.sheetnames)
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"  Sheet '{sname}': {ws.max_row} rows, {ws.max_column} columns")
        
    ws_matrix = wb['Traceability Matrix']
    print("\nSample rows from Traceability Matrix:")
    for row in list(ws_matrix.iter_rows(values_only=True))[1:6]:
        print("  ", row[:7])
        
if __name__ == '__main__':
    verify_wb('c:/Users/pm89542/Desktop/parse/Sample_Requirements_and_Test_Cases.xlsx')
