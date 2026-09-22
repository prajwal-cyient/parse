import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

class ExcelExporter:
    """
    Exports DO-178C Level B test case suites to publication-grade Excel workbooks.
    Formats headers, wrapped text, column widths, and gridlines.
    """

    @staticmethod
    def export_testcases(test_cases, output_path, title_text="DO-178C LEVEL B VERIFICATION TEST CASES"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Verification Test Cases"
        ws.views.sheetView[0].showGridLines = True

        # Title Banner
        ws.merge_cells("A1:K1")
        t_cell = ws["A1"]
        t_cell.value = f"{title_text.upper()} ({len(test_cases)} TOTAL)"
        t_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        t_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Dark Navy Blue
        t_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        # Column Headers
        headers = [
            "#", "Test Case ID", "Requirement Trace", "Test Type", 
            "Test Case Description", "Initial Condition(s)", "Test Inputs", 
            "Expected Result(s)", "Pass / Fail Criteria", "Related Requirements", 
            "Test Procedure Notes"
        ]
        ws.row_dimensions[2].height = 28
        h_fill = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid")
        h_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC")
        )

        for c_idx, h_t in enumerate(headers, 1):
            cell = ws.cell(row=2, column=c_idx, value=h_t)
            cell.fill = h_fill
            cell.font = h_font
            cell.alignment = h_align
            cell.border = thin_border

        # Populate Test Cases
        for r_idx, tc in enumerate(test_cases, 3):
            ws.cell(row=r_idx, column=1, value=r_idx-2).alignment = Alignment(vertical="top", horizontal="center")
            ws.cell(row=r_idx, column=2, value=tc.get("test_case_id", "")).alignment = Alignment(vertical="top")
            ws.cell(row=r_idx, column=3, value=tc.get("requirement_id", "")).alignment = Alignment(vertical="top")
            ws.cell(row=r_idx, column=4, value=tc.get("test_type", "NORMAL")).alignment = Alignment(vertical="top", horizontal="center")
            ws.cell(row=r_idx, column=5, value=tc.get("description", "")).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r_idx, column=6, value=tc.get("initial_condition", "")).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r_idx, column=7, value=tc.get("test_inputs", "")).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r_idx, column=8, value=tc.get("expected_result", "")).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r_idx, column=9, value=tc.get("pass_criteria", "")).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=r_idx, column=10, value=tc.get("related_requirements", "")).alignment = Alignment(vertical="top")
            ws.cell(row=r_idx, column=11, value=tc.get("test_procedure_notes", "")).alignment = Alignment(vertical="top", wrap_text=True)
            
            for c in range(1, 12):
                ws.cell(row=r_idx, column=c).border = thin_border
            ws.row_dimensions[r_idx].height = 42

        # Column Widths
        col_widths = {1: 8, 2: 28, 3: 18, 4: 14, 5: 45, 6: 35, 7: 45, 8: 45, 9: 35, 10: 16, 11: 30}
        for c_i, w in col_widths.items():
            ws.column_dimensions[get_column_letter(c_i)].width = w

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        try:
            wb.save(output_path)
            return output_path
        except PermissionError:
            import time
            base, ext = os.path.splitext(output_path)
            fallback_path = f"{base}_{int(time.time())}{ext}"
            wb.save(fallback_path)
            return fallback_path
