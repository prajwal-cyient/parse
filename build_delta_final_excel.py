import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from copy import copy
import os

def build_delta_final():
    src_path = r"c:\Users\pm89542\Desktop\parse\SW_Requirements_Sample_1_Test_Cases_DELTA_APPROVED.xlsx"
    old_baseline_path = r"c:\Users\pm89542\Desktop\parse\SW_Requirements_Sample_1_Test_Cases_20260804_185720.xlsx"
    out_dir = r"c:\Users\pm89542\Desktop\parse\app\outputs"
    os.makedirs(out_dir, exist_ok=True)
    out_path_1 = os.path.join(out_dir, "SW_Requirements_Sample_1_Test_Cases_DELTA_FINAL.xlsx")
    out_path_2 = r"c:\Users\pm89542\Desktop\parse\SW_Requirements_Sample_1_Test_Cases_DELTA_FINAL.xlsx"

    wb = openpyxl.load_workbook(src_path)
    ws = wb.active

    yellow_fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
    regular_font = Font(name="Calibri", size=10, color="000000")
    bold_font = Font(name="Calibri", size=10, bold=True, color="000000")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # 1. OBS 5: Physically DELETE rows 109 to 116 (8 tombstone rows for LRUSWRS-1000)
    # Note: openpyxl delete_rows(109, 8) deletes rows 109 through 116 inclusive.
    ws.delete_rows(109, 8)

    # 2. OBS 1: Replace LRUSWRS-1011 single row with 21 standalone rows
    # After deleting 8 rows above, the original Row 148 is now at Row 140 (148 - 8 = 140).
    target_1011_row = None
    for r in range(1, ws.max_row + 1):
        req = str(ws.cell(r, 3).value or "")
        tc = str(ws.cell(r, 2).value or "")
        if "1011" in req or "1011" in tc:
            target_1011_row = r
            break

    print(f"Target row for LRUSWRS-1011 after deletion: {target_1011_row}")

    # Delete the single combined 1011 row
    ws.delete_rows(target_1011_row, 1)

    # Insert 21 blank rows at target_1011_row
    ws.insert_rows(target_1011_row, 21)

    # Populate 21 combinations
    atypes = ["ATYPE_1", "ATYPE_2", "ATYPE_3"]
    lrus = [f"LRU_{i}" for i in range(1, 8)]

    combinations = []
    tc_counter = 1
    for atype in atypes:
        for lru in lrus:
            tc_id = f"SWVCP_AAP_TC_1011_PSR_{tc_counter:02d}"
            combinations.append((tc_id, atype, lru))
            tc_counter += 1

    for idx, (tc_id, atype, lru) in enumerate(combinations):
        r = target_1011_row + idx
        ws.cell(row=r, column=1, value="")  # Sl No will be re-indexed below
        ws.cell(row=r, column=2, value=tc_id).fill = yellow_fill
        ws.cell(row=r, column=3, value="LRUSWRS-1011")
        ws.cell(row=r, column=4, value="NORMAL")
        ws.cell(row=r, column=5, value=f"Evaluates hardware pin-strapping configuration for {atype} and {lru} per Table 1011 to ensure correct asset type and LRU identification.").fill = yellow_fill
        ws.cell(row=r, column=6, value=f"System initialized in normal operating state; Target Asset Type: {atype}; Target LRU: {lru}; Hardware pin strapping configured for {atype}_{lru}.").fill = yellow_fill
        ws.cell(row=r, column=7, value=f"Hardware Pin Strapping Inputs (ID0_DSP..ID7_DSP) driven for {atype} and {lru} configuration with valid parity.").fill = yellow_fill
        ws.cell(row=r, column=8, value=f"Decoded Asset Type = {atype}; Decoded LRU Number = {lru}; Decode_Valid = True.").fill = yellow_fill
        ws.cell(row=r, column=9, value="Observed software behaviour matches expected_result.")
        ws.cell(row=r, column=10, value="None")
        ws.cell(row=r, column=11, value="Verification mechanism to be defined by the verification environment.")

        # Apply fonts and borders
        for col in range(1, 12):
            cell = ws.cell(row=r, column=col)
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # 3. Re-index Column 1 (Sl No / Index) for all data rows
    sl_no = 1
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 2).value:
            ws.cell(r, 1, value=sl_no)
            sl_no += 1

    # Save to output locations
    wb.save(out_path_1)
    wb.save(out_path_2)
    print(f"Successfully saved DELTA_FINAL Excel to {out_path_1} and {out_path_2}")

if __name__ == "__main__":
    build_delta_final()
