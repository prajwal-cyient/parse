import json
import openpyxl
import os

def sync_348_to_all():
    excel_final_path = r"c:\Users\pm89542\Desktop\parse\SW_Requirements_Sample_1_Test_Cases_DELTA_FINAL.xlsx"
    ks_path = r"c:\Users\pm89542\Desktop\parse\app\backend\knowledge_suite.json"
    cache_path = r"c:\Users\pm89542\Desktop\parse\app\cache\cache_SW_Requirements_Sample_1.json"
    ui_out_json = r"c:\Users\pm89542\Desktop\parse\app\ui_outputs\SW_Requirements_Sample_1_generated_testcases.json"
    ui_out_excel = r"c:\Users\pm89542\Desktop\parse\app\ui_outputs\SW_Requirements_Sample_1_Test_Cases.xlsx"

    wb = openpyxl.load_workbook(excel_final_path)
    ws = wb.active

    tc_list = []
    tc_map = {}

    for r in range(3, ws.max_row + 1):
        sl_no = ws.cell(r, 1).value
        tc_id = str(ws.cell(r, 2).value or "").strip()
        req_id = str(ws.cell(r, 3).value or "").strip()
        test_type = str(ws.cell(r, 4).value or "NORMAL").strip()
        desc = str(ws.cell(r, 5).value or "").strip()
        init = str(ws.cell(r, 6).value or "").strip()
        inp = str(ws.cell(r, 7).value or "").strip()
        exp = str(ws.cell(r, 8).value or "").strip()
        criteria = str(ws.cell(r, 9).value or "").strip()
        comments = str(ws.cell(r, 10).value or "").strip()
        notes = str(ws.cell(r, 11).value or "").strip()

        if tc_id or req_id:
            tc_dict = {
                "test_case_id": tc_id,
                "requirement_id": req_id,
                "test_type": test_type,
                "description": desc,
                "initial_condition": init,
                "test_inputs": inp,
                "expected_result": exp,
                "pass_criteria": criteria,
                "related_requirements": comments if comments != "None" else "",
                "test_procedure_notes": notes
            }
            tc_list.append(tc_dict)
            tc_map.setdefault(req_id, []).append(tc_dict)

    # Write to all JSON targets
    with open(ks_path, "w", encoding="utf-8") as f:
        json.dump(tc_map, f, indent=2)

    with open(cache_path, "w", encoding="utf-8") as f:
        json.dump(tc_map, f, indent=2)

    with open(ui_out_json, "w", encoding="utf-8") as f:
        json.dump(tc_list, f, indent=2)

    wb.save(ui_out_excel)

    print(f"Sync complete! Exactly {len(tc_list)} test cases written across all cache and output files.")

if __name__ == "__main__":
    sync_348_to_all()
