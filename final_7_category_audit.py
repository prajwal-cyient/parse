import openpyxl, json, re
from collections import Counter

wb = openpyxl.load_workbook('app/outputs/SW_Requirements_Sample_1_Test_Cases_PERFECTED_APPROVED.xlsx')
ws = wb.active
headers = [c for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
tcs = []
for row in ws.iter_rows(min_row=3, values_only=True):
    tcs.append(dict(zip(headers, row)))

print(f"=== FULL AUDIT REPORT: {len(tcs)} TEST CASES ===")

# Category 1: Internal Contradictions
cat1_issues = []
for tc in tcs:
    inp = str(tc['Test Inputs'] or '').lower()
    er = str(tc['Expected Result(s)'] or '').lower()
    pc = str(tc['Pass / Fail Criteria'] or '').lower()
    desc = str(tc['Test Case Description'] or '').lower()
    
    if 'true and false' in inp or 'false and true' in inp or 'true and false' in er:
        cat1_issues.append((tc['Test Case ID'], 'Boolean contradiction'))
    if 'harmonize_failed_stl = true' in er and 'harmonizing_failed = false' in er:
        cat1_issues.append((tc['Test Case ID'], 'Contradicting failure flags in ER'))

# Category 2: Wrong Data Types
cat2_issues = []
for tc in tcs:
    for f in ['Test Inputs', 'Expected Result(s)', 'Pass / Fail Criteria']:
        val = str(tc[f] or '')
        for var in ['Harmonize_Offset_Default', 'Harmonize_Offset_New', 'Harmonize_SFC_Default', 'Act_Disp_Raw', 'K_Harmonize_New', 'Profile_Id']:
            m = re.search(rf'{var}\s*=\s*(True|False)', val, re.IGNORECASE)
            if m:
                cat2_issues.append((tc['Test Case ID'], f"Boolean on numeric var {var}: '{m.group(0)}' in {f}"))

# Category 3: Placeholder Input Values
cat3_issues = []
placeholder_patterns = ['per requirement', 'valid value', 'nominal value', 'as specified', 'applicable value', 'valid operational value', 'verification mechanism to be defined']
for tc in tcs:
    inp = str(tc['Test Inputs'] or '').lower()
    if any(p in inp for p in placeholder_patterns):
        cat3_issues.append((tc['Test Case ID'], inp))

# Category 4: Requirement-Level Calculation Validation
cat4_issues = []
for tc in tcs:
    er = str(tc['Expected Result(s)'] or '')
    calc_match = re.search(r'\(\s*([\d\.\-]+)\s*\*\s*([\d\.\-]+)\s*\)\s*\+\s*([\d\.\-]+)\s*=\s*([\d\.\-]+)', er)
    if calc_match:
        a, b, c, d = map(float, calc_match.groups())
        expected_calc = round((a * b) + c, 4)
        if abs(expected_calc - d) > 0.001:
            cat4_issues.append((tc['Test Case ID'], er, f"Expected {expected_calc} but got {d}"))

# Category 5: Duplicate / Redundant Test Cases
cat5_issues = []
ids = [tc['Test Case ID'] for tc in tcs]
id_counts = Counter(ids)
for tid, cnt in id_counts.items():
    if cnt > 1:
        cat5_issues.append((tid, f"Appears {cnt} times"))

# Category 6: Executable Test Procedure
cat6_issues = []
for tc in tcs:
    pn = str(tc['Test Procedure Notes'] or '').strip()
    if not pn or len(pn) < 10 or 'to be defined' in pn.lower():
        cat6_issues.append((tc['Test Case ID'], pn))

# Category 7: Pass Criteria Consistency
cat7_issues = []
for tc in tcs:
    er = str(tc['Expected Result(s)'] or '')
    pc = str(tc['Pass / Fail Criteria'] or '')
    if 'equals true' in pc.lower() and ' = true' not in er.lower() and 'is true' not in er.lower() and 'true' not in er.lower():
        cat7_issues.append((tc['Test Case ID'], f"PC has True mismatch: ER='{er}' | PC='{pc}'"))
    if 'equals false' in pc.lower() and ' = false' not in er.lower() and 'is false' not in er.lower() and 'false' not in er.lower():
        cat7_issues.append((tc['Test Case ID'], f"PC has False mismatch: ER='{er}' | PC='{pc}'"))

print(f"1. Internal Contradictions:       {len(cat1_issues)} issues")
print(f"2. Wrong Data Types:              {len(cat2_issues)} issues")
print(f"3. Placeholder Input Values:      {len(cat3_issues)} issues")
print(f"4. Calculation Mismatches:        {len(cat4_issues)} issues")
print(f"5. Duplicate Test Case IDs:       {len(cat5_issues)} issues")
print(f"6. Test Procedure Issues:         {len(cat6_issues)} issues")
print(f"7. Pass Criteria Inconsistencies: {len(cat7_issues)} issues")
print(f"Total Combined Defects:           {len(cat1_issues) + len(cat2_issues) + len(cat3_issues) + len(cat4_issues) + len(cat5_issues) + len(cat6_issues) + len(cat7_issues)}")
