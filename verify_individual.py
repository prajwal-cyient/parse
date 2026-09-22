import openpyxl
import glob
import os

for path in sorted(glob.glob('c:/Users/pm89542/Desktop/parse/sample *.xlsx')):
    wb = openpyxl.load_workbook(path)
    print(f"=== {os.path.basename(path)} ===")
    for sname in wb.sheetnames:
        ws = wb[sname]
        print(f"  Sheet '{sname}': {ws.max_row} rows, {ws.max_column} cols")
