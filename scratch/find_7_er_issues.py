import openpyxl
import json
import re

wb = openpyxl.load_workbook('app/outputs/SW_Requirements_Sample_1_Test_Cases_PERFECTED_APPROVED.xlsx')
ws = wb.active

headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
col_idx = {h: i + 1 for i, h in enumerate(headers)}

for r in range(2, ws.max_row + 1):
    tc_id = str(ws.cell(r, col_idx["Test Case ID"]).value or "")
    desc = str(ws.cell(r, col_idx["Test Description"]).value or "")
    er = str(ws.cell(r, col_idx["Expected Results"]).value or "")
    
    if "false" in desc.lower() or "outside" in desc.lower() or "boundary" in desc.lower() or "robustness" in desc.lower() or "error" in desc.lower() or "invalid" in desc.lower():
        if "= False" not in er and "= True" not in er and "False" not in er:
            print(f"Row {r} | TC: {tc_id}")
            print(f"   DESC: {desc}")
            print(f"   ER:   {er}\n")
