import openpyxl
import re

wb = openpyxl.load_workbook('app/outputs/SW_Requirements_Sample_1_Test_Cases_PERFECTED_APPROVED.xlsx')
ws = wb.active

headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
col_idx = {h: i + 1 for i, h in enumerate(headers)}

rows = []
for r in range(2, ws.max_row + 1):
    row_dict = {h: str(ws.cell(r, col_idx[h]).value or "").strip() for h in headers}
    row_dict["_row_num"] = r
    rows.append(row_dict)

print(f"Total Rows Analyzed: {len(rows)}")

# Check Observation 1:
obs1_issues = []
table_1011_rows = [r for r in rows if "1011" in r["Requirement ID"]]
if len(table_1011_rows) != 21:
    obs1_issues.append(f"Expected 21 test cases for Table 1011, found {len(table_1011_rows)}")
for r in table_1011_rows:
    ic = r["Initial Conditions"]
    tc_id = r["Test Case ID"]
    if not re.search(r'ATYPE_[123]', ic) or not re.search(r'LRU_[1-7]', ic):
        obs1_issues.append(f"TC {tc_id} missing ATYPE or LRU in Initial Conditions: '{ic}'")
    if not tc_id.startswith("SWVCP_AAP_TC_1011_"):
        obs1_issues.append(f"TC {tc_id} naming does not match SWVCP_AAP_TC_1011_<ATYPE>_<LRU>")

# Check Observation 2:
obs2_issues = []
for r in rows:
    tc_id = r["Test Case ID"]
    ic = r["Initial Conditions"]
    ti = r["Test Inputs"]
    if "FLIGHT_MODE" not in ic:
        obs2_issues.append(f"TC {tc_id} missing FLIGHT_MODE in Initial Conditions")
    if re.search(r'\bcommand(?:ed)?\s+to\b', ti, re.IGNORECASE):
        obs2_issues.append(f"TC {tc_id} has commanded position in Test Inputs: '{ti}'")

# Check Observation 3:
obs3_issues = []
forbidden_inputs = [
    r'Act_Disp_Range_\w+_Fault',
    r'Harmonizing_PSR_Data_Fault',
    r'Harmonizing_Failed',
    r'Contract_Stop_Collection_Complete',
    r'Expand_Stop_Collection_Complete',
    r'VOID_Collection_Complete',
    r'Harmonizing_Complete',
    r'Harmonizing_Sequence_Complete',
    r'Sequence_Complete'
]
for r in rows:
    tc_id = r["Test Case ID"]
    ti = r["Test Inputs"]
    for f in forbidden_inputs:
        if re.search(r'\b' + f + r'\b', ti, re.IGNORECASE):
            obs3_issues.append(f"TC {tc_id} has output flag '{f}' in Test Inputs: '{ti}'")

# Check Observation 4:
obs4_issues = []
for r in rows:
    tc_id = r["Test Case ID"]
    desc = r["Test Description"]
    er = r["Expected Results"]
    if any(k in desc.lower() for k in ["false", "outside", "invalid", "corrupt", "not performed", "does not transmit", "exceeds"]):
        if "= False" not in er and "= True" not in er and "False" not in er and "True" not in er:
            obs4_issues.append(f"TC {tc_id} missing explicit boolean assertion in ER: '{er}' (Desc: {desc})")

# Check Observation 5:
obs5_issues = []
step4_clubbed = [r for r in rows if "4a+4b" in r["Requirement ID"] or "4a" in r["Requirement ID"].lower()]
step9_clubbed = [r for r in rows if "9a+9b" in r["Requirement ID"] or "9a" in r["Requirement ID"].lower()]
step14_clubbed = [r for r in rows if "14a+14b" in r["Requirement ID"] or "14a" in r["Requirement ID"].lower()]
if not step4_clubbed:
    obs5_issues.append("Contract fault sequence (Steps 4a+4b) is not clubbed")
if not step9_clubbed:
    obs5_issues.append("Expand fault sequence (Steps 9a+9b) is not clubbed")
if not step14_clubbed:
    obs5_issues.append("VOID fault sequence (Steps 14a+14b) is not clubbed")

# Check Observation 6:
obs6_issues = []
psr_1000 = [r for r in rows if r["Requirement ID"] == "LRUSWRS-1000"]
if len(psr_1000) != 1:
    obs6_issues.append(f"Expected exactly 1 consolidated PSR 1000 test case, found {len(psr_1000)}")
else:
    er = psr_1000[0]["Expected Results"]
    ti = psr_1000[0]["Test Inputs"]
    if "IVT_Mode_ModeLgc = True" not in ti or "Harmonizing_Active_STL = True" not in ti:
        obs6_issues.append(f"PSR 1000 Test Inputs missing standard stimuli: '{ti}'")
    for m in ["Harmonize_SFC_PSR", "Harmonize_Offset_PSR", "Config_LRU_PSR", "Asset_ID_PSR", "Upper_Limit_LRU_PSR", "Lower_Limit_LRU_PSR", "CRC_PSR", "Var_ID_PSR", "Label_Version_PSR"]:
        if m not in er:
            obs6_issues.append(f"PSR 1000 missing member '{m}' in Expected Results")

# Check Observation 7:
obs7_issues = []
for r in rows:
    tc_id = r["Test Case ID"]
    ti = r["Test Inputs"]
    if ti.lower() in ["none", "none.", "n/a", "", "."]:
        obs7_issues.append(f"TC {tc_id} has empty/None Test Inputs")
    if "generic" in ti.lower() or "placeholder" in ti.lower() or "tbd" in ti.lower():
        obs7_issues.append(f"TC {tc_id} has placeholder in Test Inputs: '{ti}'")

# Check Observation 8:
obs8_issues = []
for r in rows:
    tc_id = r["Test Case ID"]
    ic = r["Initial Conditions"]
    ti = r["Test Inputs"]
    if re.search(r'\bIVT_Mode_ModeLgc\s*=\s*(?:True|False)\b', ic, re.IGNORECASE):
        obs8_issues.append(f"TC {tc_id} has dynamic flag 'IVT_Mode_ModeLgc' in Initial Conditions: '{ic}'")

print("\n--- RESULTS FOR THE 8 CLIENT OBSERVATIONS ---")
observations = [
    ("Observation #1: LRU_1 TO LRU_7 x ATYPE_1 TO ATYPE_3 COMBINATIONS (Table 1011)", obs1_issues),
    ("Observation #2: CLUB RELATED REQUIREMENTS & DEFINE COMPLETE INITIAL CONDITIONS", obs2_issues),
    ("Observation #3: STRICT THREE-TIER SEPARATION FOR INPUTS, OUTPUTS, AND SETUP", obs3_issues),
    ("Observation #4: EXPLICIT BOOLEAN NEGATIVE ASSERTIONS", obs4_issues),
    ("Observation #5: SEQUENCE STEP CLUBBING & ZERO FRAGMENTATION", obs5_issues),
    ("Observation #6: CONSOLIDATED MULTI-MEMBER DATA RECORDS (PSR 1000)", obs6_issues),
    ("Observation #7: VALID OPERATIONAL TEST INPUTS (No 'None' or Placeholders)", obs7_issues),
    ("Observation #8: DYNAMIC MODE FLAG IN TEST INPUTS, NOT INITIAL CONDITIONS", obs8_issues),
]

for title, issues in observations:
    status = "PASS" if len(issues) == 0 else ("PARTIAL" if len(issues) < 5 else "FAIL")
    print(f"\n[{status}] {title}")
    if issues:
        print(f"   Issues found ({len(issues)}):")
        for iss in issues[:5]:
            print(f"     * {iss}")
