import sys, os
sys.path.insert(0, 'app/backend')
from docx_parser import UniversalDocxParser
import json

parsed = UniversalDocxParser.parse_document('SW_Requirements_Sample_4.docx')
print('Sample 4 total requirements found:', parsed['total_requirements_found'])
for i, req in enumerate(parsed['requirements']):
    print(f"Req {i+1}: ID={req.get('id')} | Type={req.get('type')}")
    print(f"   Text preview: {req.get('text')[:120]}...\n")

# Check if there is cache or output files for Sample 4
cache_path = 'app/cache/cache_SW_Requirements_Sample_4.json'
out_path = 'app/ui_outputs/SW_Requirements_Sample_4_generated_testcases.json'
print(f"Cache exists ({cache_path}):", os.path.exists(cache_path))
if os.path.exists(cache_path):
    with open(cache_path, 'r', encoding='utf-8') as f:
        cdata = json.load(f)
        print(f"Cache items count: {len(cdata)}")
print(f"Output exists ({out_path}):", os.path.exists(out_path))
if os.path.exists(out_path):
    with open(out_path, 'r', encoding='utf-8') as f:
        odata = json.load(f)
        print(f"Output items count: {len(odata)}")

# Check sample 4.xlsx
import openpyxl
if os.path.exists('sample 4.xlsx'):
    wb = openpyxl.load_workbook('sample 4.xlsx')
    ws = wb['Test Cases Catalog'] if 'Test Cases Catalog' in wb.sheetnames else wb.active
    print(f"sample 4.xlsx rows in '{ws.title}': {ws.max_row}")
