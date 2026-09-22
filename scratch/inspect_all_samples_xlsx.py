import openpyxl

for fname in ['sample 1.xlsx', 'sample 2.xlsx', 'sample 3.xlsx', 'sample 4.xlsx']:
    wb = openpyxl.load_workbook(fname)
    sheet = 'Test Cases Catalog' if 'Test Cases Catalog' in wb.sheetnames else wb.active.title
    ws = wb[sheet]
    print(f"File: {fname} -> Sheet: {sheet} -> Total TCs: {ws.max_row - 1}")
    for r in range(2, min(ws.max_row + 1, 6)):
        tc_id = ws.cell(r, 2).value or ws.cell(r, 1).value
        req_id = ws.cell(r, 1).value or ws.cell(r, 2).value
        print(f"   Row {r-1}: TC={tc_id} | REQ={req_id}")
