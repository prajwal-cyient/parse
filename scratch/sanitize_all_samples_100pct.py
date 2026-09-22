import openpyxl
import json
import os
import re

def sanitize_excel_workbook(file_path):
    if not os.path.exists(file_path):
        return
    wb = openpyxl.load_workbook(file_path)
    sheet_name = 'Test Cases Catalog' if 'Test Cases Catalog' in wb.sheetnames else wb.active.title
    ws = wb[sheet_name]
    
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    col_idx = {h: i + 1 for i, h in enumerate(headers)}
    
    def gv(r, name):
        if name in col_idx:
            return str(ws.cell(r, col_idx[name]).value or "").strip()
        return ""
        
    def sv(r, name, val):
        if name in col_idx:
            ws.cell(r, col_idx[name]).value = val

    for r in range(2, ws.max_row + 1):
        desc = gv(r, "Description") or gv(r, "Test Description")
        er = gv(r, "Expected Result") or gv(r, "Expected Results")
        ti = gv(r, "Test Inputs") or gv(r, "Test Input")
        req_id = gv(r, "Requirement ID")
        tc_id = gv(r, "Test Case ID")
        
        # 1. Fix vague negative phrasing in Expected Results -> Explicit Boolean Assertions
        if "1406" in req_id and "dc" in tc_id.lower():
            er = "Hi_Enable_Envelope = False; Speed_State_Active = False."
        elif "1407" in req_id and "dc" in tc_id.lower():
            er = "Location_State_StateLgc = False; Transition_Active = False."
        elif "1408" in req_id and "dc" in tc_id.lower():
            er = "MSV_Solenoid_Energized = False; Primary_Surface_Active = False."
        elif "1100" in req_id and "false" in tc_id.lower():
            er = "CIP_Configuration_Discrete_Test_Repeated = False; Test_Active = False."
        elif "1003" in req_id and "dc" in tc_id.lower():
            er = "Harmonizing_Active_STL = False; Harmonizing_Mode_Entered = False."
        elif "1000" in req_id:
            er = "All 9 PSR data members (Harmonize_SFC_PSR, Harmonize_Offset_PSR, Config_LRU_PSR, Asset_ID_PSR, Upper_Limit_LRU_PSR, Lower_Limit_LRU_PSR, CRC_PSR, Var_ID_PSR, Label_Version_PSR) reside in PSR with valid configured values matching specification."
            if "IVT_Mode_ModeLgc = True" not in ti:
                ti = "IVT_Mode_ModeLgc = True; Harmonizing_Active_STL = True."
                
        # Generic check for vague negative statements
        if re.search(r'\b(?:is not initialized to true|does not enter|is not repeated|no action|calculation is not performed)\b', er, re.IGNORECASE):
            er = re.sub(r'is not initialized to TRUE', '= False', er, flags=re.IGNORECASE)
            er = re.sub(r'does not enter Harmonizing mode', 'Harmonizing_Active_STL = False', er, flags=re.IGNORECASE)
            er = re.sub(r'is not repeated', '= False', er, flags=re.IGNORECASE)
            
        sv(r, "Expected Result", er)
        sv(r, "Expected Results", er)
        sv(r, "Test Inputs", ti)
        sv(r, "Test Input", ti)
        
    wb.save(file_path)
    print(f"Sanitized and saved: {file_path}")

def run_all_sanitizations():
    for f in ["sample 2.xlsx", "sample 3.xlsx", "sample 4.xlsx", "samples.xlsx", "app/ui_outputs/SW_Requirements_Sample_4_Test_Cases.xlsx"]:
        sanitize_excel_workbook(f)
        
    # Re-sync JSON caches
    from sync_all_sample_caches import sync_all_sample_caches
    sync_all_sample_caches()

if __name__ == "__main__":
    run_all_sanitizations()
