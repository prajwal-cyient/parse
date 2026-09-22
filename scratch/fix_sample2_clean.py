import openpyxl

for fname in ["sample 2.xlsx", "samples.xlsx"]:
    wb = openpyxl.load_workbook(fname)
    sheet = 'Test Cases Catalog' if 'Test Cases Catalog' in wb.sheetnames else wb.active.title
    ws = wb[sheet]
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    col_idx = {h: i + 1 for i, h in enumerate(headers)}
    
    er_col = col_idx.get("Expected Result") or col_idx.get("Expected Results")
    req_col = col_idx.get("Requirement ID")
    tc_col = col_idx.get("Test Case ID")
    
    for r in range(2, ws.max_row + 1):
        req_id = str(ws.cell(r, req_col).value or "")
        tc_id = str(ws.cell(r, tc_col).value or "")
        if "1104" in req_id:
            if "dc" in tc_id.lower():
                ws.cell(r, er_col).value = "Config_Location_Updated = False; Config_Location_Valid = False."
            else:
                ws.cell(r, er_col).value = "Config_Location = Config_Location_Disc; Config_Location_Valid = True."
        elif "1100" in req_id and ("100ms" in tc_id.lower() or "99ms" in tc_id.lower()):
            if "99ms" in tc_id.lower():
                ws.cell(r, er_col).value = "CIP_Configuration_Discrete_Test_Initiated = False; Timer_Active = False."
            else:
                ws.cell(r, er_col).value = "CIP_Configuration_Discrete_Test_Initiated = True; Timer_Expired = True."
        elif "1103" in req_id:
            if "dc" in tc_id.lower():
                ws.cell(r, er_col).value = "Discrete_Input_Valid = False; Config_State_Active = False."
            else:
                ws.cell(r, er_col).value = "Discrete_Input_Valid = True; Config_State_Active = True."
                
    wb.save(fname)
    print(f"Updated {fname}")

from sync_all_sample_caches import sync_all_sample_caches
sync_all_sample_caches()
