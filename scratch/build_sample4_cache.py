import openpyxl
import json
import os

def load_sample4():
    wb = openpyxl.load_workbook('sample 4.xlsx')
    ws = wb['Test Cases Catalog']
    
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    col_idx = {h: i + 1 for i, h in enumerate(headers)}
    
    test_cases = []
    for r in range(2, ws.max_row + 1):
        def gv(name):
            if name in col_idx:
                return str(ws.cell(r, col_idx[name]).value or "").strip()
            return ""
            
        tc_id = gv("Test Case ID")
        req_id = gv("Requirement ID")
        if not tc_id or not req_id:
            continue
            
        desc = gv("Description")
        test_type = gv("Test Type") or "NORMAL"
        ic = gv("Initial Condition")
        ti = gv("Test Inputs")
        er = gv("Expected Result")
        pc = gv("Pass Criteria")
        rel = gv("Related Requirements")
        notes = gv("Test Procedure Notes")
        
        tc = {
            "test_case_id": tc_id,
            "requirement_id": req_id,
            "test_type": test_type,
            "description": desc,
            "initial_condition": ic,
            "test_inputs": ti,
            "expected_result": er,
            "pass_criteria": pc or f"Observed outputs equal expected result: {er[:60]}.",
            "related_requirements": rel,
            "test_procedure_notes": notes or "Verify via internal registers and STL_Bus telemetry."
        }
        test_cases.append(tc)
        
    print(f"Total test cases loaded from sample 4.xlsx: {len(test_cases)}")
    
    cache_path = 'app/cache/cache_SW_Requirements_Sample_4.json'
    ui_out_path = 'app/ui_outputs/SW_Requirements_Sample_4_generated_testcases.json'
    
    with open(cache_path, 'w', encoding='utf-8') as f:
        json.dump(test_cases, f, indent=2)
    print(f"Saved {len(test_cases)} test cases to {cache_path}")
    
    with open(ui_out_path, 'w', encoding='utf-8') as f:
        json.dump(test_cases, f, indent=2)
    print(f"Saved {len(test_cases)} test cases to {ui_out_path}")

    # Also export Excel for Sample 4
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "DO-178C Verification Test Cases"
    excel_headers = ["Test Case ID", "Requirement ID", "Test Description", "Initial Conditions", "Test Inputs", "Expected Results", "Pass/Fail Criteria", "Related Requirements", "Test Procedure Notes"]
    ws_out.append(excel_headers)
    for tc in test_cases:
        ws_out.append([
            tc.get("test_case_id", ""),
            tc.get("requirement_id", ""),
            tc.get("description", ""),
            tc.get("initial_condition", ""),
            tc.get("test_inputs", ""),
            tc.get("expected_result", ""),
            tc.get("pass_criteria", ""),
            tc.get("related_requirements", ""),
            tc.get("test_procedure_notes", "")
        ])
    excel_path = "app/ui_outputs/SW_Requirements_Sample_4_Test_Cases.xlsx"
    wb_out.save(excel_path)
    print(f"Saved Excel to {excel_path}")

if __name__ == "__main__":
    load_sample4()
