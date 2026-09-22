import openpyxl
import json

wb = openpyxl.load_workbook('sample 4.xlsx')
ws = wb['Test Cases Catalog']

print(f"--- SAMPLE 4.XLSX TEST CASES ({ws.max_row - 1} rows) ---")
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
print("Headers:", headers)

col_idx = {h: i + 1 for i, h in enumerate(headers)}

sample4_reference_tcs = []
for r in range(2, ws.max_row + 1):
    tc = {h: str(ws.cell(r, col_idx[h]).value or "").strip() for h in headers}
    sample4_reference_tcs.append(tc)
    print(f"[{r-1}] TC_ID: {tc.get('Test Case ID')} | REQ: {tc.get('Requirement ID')} | Type: {tc.get('Test Type') or tc.get('Test Description')[:30]}")

print(f"\nTotal Reference Test Cases in sample 4.xlsx: {len(sample4_reference_tcs)}")
