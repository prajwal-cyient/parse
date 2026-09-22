import os
import re
import json
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def find_json_objects(text):
    results = []
    i = 0
    n = len(text)
    while i < n:
        if text[i] in '{[':
            start = i
            stack = [text[i]]
            i += 1
            in_string = False
            escape = False
            while i < n and stack:
                ch = text[i]
                if escape:
                    escape = False
                elif ch == '\\' and in_string:
                    escape = True
                elif ch == '"':
                    in_string = not in_string
                elif not in_string:
                    if ch in '{[':
                        stack.append(ch)
                    elif ch == '}' and stack[-1] == '{':
                        stack.pop()
                    elif ch == ']' and stack[-1] == '[':
                        stack.pop()
                i += 1
            if not stack:
                candidate = text[start:i]
                try:
                    obj = json.loads(candidate, strict=False)
                    results.append((start, i, obj))
                except Exception:
                    cleaned = re.sub(r'[\r\n]+', ' ', candidate)
                    try:
                        obj = json.loads(cleaned, strict=False)
                        results.append((start, i, obj))
                    except Exception:
                        pass
        else:
            i += 1
    return results

def parse_file_content(filename, content):
    pattern = re.compile(r'(?:\r?\n|^)\s*(\d+)\s*\.\s*requirement\b', re.IGNORECASE)
    matches = list(pattern.finditer(content))
    
    parsed_reqs = []
    
    for idx, m in enumerate(matches):
        req_num = m.group(1)
        start_pos = m.start()
        end_pos = matches[idx + 1].start() if idx + 1 < len(matches) else len(content)
        block = content[start_pos:end_pos]
        
        # 1. ID
        id_match = re.search(r'ID\s*:\s*([^\r\n]+)', block)
        req_id = id_match.group(1).strip() if id_match else 'N/A'
        
        # 2. Req Type
        type_match = re.search(r'Req\s+Type\s*:\s*([^\r\n]+)', block, re.IGNORECASE)
        req_type = type_match.group(1).strip() if type_match else 'N/A'
        
        # 3. Safety Impact
        safety_impact_match = re.search(r'Safety\s+Impact\s*:\s*([^\r\n]+)', block, re.IGNORECASE)
        safety_impact = safety_impact_match.group(1).strip() if safety_impact_match else 'N/A'
        
        # 4. Safety Rationale
        safety_rat_match = re.search(r'Safety\s+Rationale\s*:\s*([^\r\n]+)', block, re.IGNORECASE)
        safety_rat = safety_rat_match.group(1).strip() if safety_rat_match else 'N/A'
        
        # 5. Extract JSON test cases & ranges
        json_objs = find_json_objects(block)
        test_cases = []
        json_ranges = []
        for j_start, j_end, obj in json_objs:
            json_ranges.append((j_start, j_end))
            if isinstance(obj, dict) and 'test_cases' in obj:
                test_cases.extend(obj['test_cases'])
            elif isinstance(obj, list):
                test_cases.extend(obj)
                
        # Slice out JSON substrings from block to leave only pure text & metadata
        non_json_block = ""
        last_idx = 0
        for j_start, j_end in sorted(json_ranges):
            non_json_block += block[last_idx:j_start]
            last_idx = j_end
        non_json_block += block[last_idx:]
        
        # 6. Extract Requirement Text and Notes
        lines = non_json_block.splitlines()
        clean_lines = []
        notes = []
        
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            if re.match(r'^\d+\s*\.\s*requirement\b', stripped, re.IGNORECASE) or stripped.lower() == 'requirement:':
                continue
            if re.match(r'^ID\s*:', stripped, re.IGNORECASE):
                continue
            if re.match(r'^Req\s+Type\s*:', stripped, re.IGNORECASE):
                continue
            if re.match(r'^Safety\s+Impact\s*:', stripped, re.IGNORECASE):
                continue
            if re.match(r'^Safety\s+Rationale\s*:', stripped, re.IGNORECASE):
                continue
            if stripped.startswith('```'):
                continue
            if re.match(r'^(NOTE|Note)\s*:', stripped, re.IGNORECASE):
                notes.append(stripped)
                continue
            clean_lines.append(line)
            
        req_text = "\n".join(clean_lines).strip()
        notes_text = "\n".join(notes).strip()
        
        parsed_reqs.append({
            'source_file': filename,
            'req_num': req_num,
            'req_id': req_id,
            'req_text': req_text,
            'notes': notes_text,
            'req_type': req_type,
            'safety_impact': safety_impact,
            'safety_rationale': safety_rat,
            'test_cases': test_cases
        })
        
    return parsed_reqs

def build_excel_for_file(sample_filename, file_reqs, output_file_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # remove default sheet

    font_family = 'Segoe UI'
    
    title_font = Font(name=font_family, size=15, bold=True, color='1F4E79')
    subtitle_font = Font(name=font_family, size=10, italic=True, color='595959')
    card_title_font = Font(name=font_family, size=11, bold=True, color='1F4E79')
    card_val_font = Font(name=font_family, size=18, bold=True, color='000000')
    header_font = Font(name=font_family, size=11, bold=True, color='FFFFFF')
    sub_header_font = Font(name=font_family, size=11, bold=True, color='1F4E79')
    cell_font = Font(name=font_family, size=10)
    bold_cell_font = Font(name=font_family, size=10, bold=True)
    
    navy_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    slate_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
    light_blue_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    card_fill = PatternFill(start_color='F2F4F8', end_color='F2F4F8', fill_type='solid')
    alt_row_fill = PatternFill(start_color='F9FAFC', end_color='F9FAFC', fill_type='solid')
    white_fill = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    
    safety_yes_fill = PatternFill(start_color='FADBD8', end_color='FADBD8', fill_type='solid')
    safety_no_fill = PatternFill(start_color='D4EFDF', end_color='D4EFDF', fill_type='solid')
    
    type_normal_fill = PatternFill(start_color='E8F8F5', end_color='E8F8F5', fill_type='solid')
    type_dc_fill = PatternFill(start_color='FEF9E7', end_color='FEF9E7', fill_type='solid')
    type_boundary_fill = PatternFill(start_color='F4ECF7', end_color='F4ECF7', fill_type='solid')
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    card_border = Border(
        left=Side(style='medium', color='1F4E79'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    align_left_center = Alignment(horizontal='left', vertical='center')

    # ----------------------------------------------------
    # SHEET 1: Summary Dashboard
    # ----------------------------------------------------
    ws_dash = wb.create_sheet(title="Summary Dashboard")
    ws_dash.views.sheetView[0].showGridLines = True
    
    ws_dash['A1'] = f"Requirements & Test Cases Analysis - {sample_filename}"
    ws_dash['A1'].font = title_font
    
    all_tcs = [tc for r in file_reqs for tc in r['test_cases']]
    safety_reqs = sum(1 for r in file_reqs if r['safety_impact'].lower() == 'yes')
    
    ws_dash['A2'] = f"Source File: {sample_filename} | Total Requirements: {len(file_reqs)} | Total Test Cases: {len(all_tcs)} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws_dash['A2'].font = subtitle_font
    
    # KPI Cards (Row 4-5)
    kpis = [
        ("Requirements Count", len(file_reqs), "B"),
        ("Test Cases Count", len(all_tcs), "D"),
        ("Safety Critical Reqs", safety_reqs, "F")
    ]
    
    for title, val, col in kpis:
        c1 = f"{col}4"
        c2 = f"{col}5"
        ws_dash[c1] = title
        ws_dash[c1].font = card_title_font
        ws_dash[c1].alignment = align_center
        ws_dash[c1].fill = light_blue_fill
        ws_dash[c1].border = thin_border
        
        ws_dash[c2] = val
        ws_dash[c2].font = card_val_font
        ws_dash[c2].alignment = align_center
        ws_dash[c2].fill = card_fill
        ws_dash[c2].border = card_border

    # Test Case Type Breakdown Table (Row 7)
    ws_dash['B7'] = "Test Case Type Distribution"
    ws_dash['B7'].font = Font(name=font_family, size=12, bold=True, color='1F4E79')
    
    headers_t2 = ["Test Type", "Description", "Count", "Percentage"]
    for i, h in enumerate(headers_t2, start=2):
        cell = ws_dash.cell(row=8, column=i, value=h)
        cell.font = header_font
        cell.fill = slate_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    tc_types = {}
    for tc in all_tcs:
        tt = tc.get('test_type', 'UNKNOWN')
        tc_types[tt] = tc_types.get(tt, 0) + 1
        
    type_descriptions = {
        'NORMAL': 'Standard nominal operating conditions verification',
        'DC': 'Decision Coverage / Disabling Condition verification',
        'BOUNDARY': 'Boundary value analysis and edge condition testing'
    }
    
    curr_row = 9
    total_tcs_cnt = len(all_tcs)
    for tt, count in sorted(tc_types.items()):
        pct = f"{(count / total_tcs_cnt * 100):.1f}%" if total_tcs_cnt > 0 else "0%"
        ws_dash.cell(row=curr_row, column=2, value=tt).alignment = align_center
        ws_dash.cell(row=curr_row, column=3, value=type_descriptions.get(tt, 'Specific test category')).alignment = align_left_center
        ws_dash.cell(row=curr_row, column=4, value=count).alignment = align_center
        ws_dash.cell(row=curr_row, column=5, value=pct).alignment = align_center
        
        fill = alt_row_fill if curr_row % 2 == 0 else white_fill
        for c in range(2, 6):
            ws_dash.cell(row=curr_row, column=c).font = cell_font
            ws_dash.cell(row=curr_row, column=c).fill = fill
            ws_dash.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

    # ----------------------------------------------------
    # SHEET 2: Traceability Matrix
    # ----------------------------------------------------
    ws_matrix = wb.create_sheet(title="Traceability Matrix")
    ws_matrix.views.sheetView[0].showGridLines = True
    ws_matrix.freeze_panes = 'A2'
    
    matrix_headers = [
        "Req #", "Requirement ID", "Requirement Text", "Req Type", 
        "Safety Impact", "Test Case ID", "Test Type", "Test Case Description", 
        "Initial Condition", "Test Inputs", "Expected Result", "Pass Criteria", 
        "Related Requirements", "Test Procedure Notes"
    ]
    
    for c_idx, h in enumerate(matrix_headers, start=1):
        cell = ws_matrix.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    m_row = 2
    for r in file_reqs:
        tcs = r['test_cases']
        fill = alt_row_fill if m_row % 2 == 0 else white_fill
        if not tcs:
            row_vals = [
                r['req_num'], r['req_id'], r['req_text'], r['req_type'],
                r['safety_impact'], "N/A", "N/A", "No test cases specified",
                "-", "-", "-", "-", "-", "-"
            ]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_matrix.cell(row=m_row, column=c_idx, value=val)
                cell.font = cell_font
                cell.fill = fill
                cell.alignment = align_left
                cell.border = thin_border
            m_row += 1
        else:
            for tc in tcs:
                rel_reqs = tc.get('related_requirements', 'None')
                if isinstance(rel_reqs, list):
                    rel_reqs = ", ".join(rel_reqs) if rel_reqs else "None"
                    
                row_vals = [
                    r['req_num'],
                    r['req_id'],
                    r['req_text'],
                    r['req_type'],
                    r['safety_impact'],
                    tc.get('test_case_id', 'N/A'),
                    tc.get('test_type', 'N/A'),
                    tc.get('description', ''),
                    tc.get('initial_condition', ''),
                    tc.get('test_inputs', ''),
                    tc.get('expected_result', ''),
                    tc.get('pass_criteria', ''),
                    rel_reqs,
                    tc.get('test_procedure_notes', '')
                ]
                for c_idx, val in enumerate(row_vals, start=1):
                    cell = ws_matrix.cell(row=m_row, column=c_idx, value=val)
                    cell.font = cell_font
                    cell.fill = fill
                    cell.alignment = align_left
                    cell.border = thin_border
                    
                    if c_idx == 5: # Safety Impact
                        cell.alignment = align_center
                        if str(val).strip().lower() == 'yes':
                            cell.fill = safety_yes_fill
                            cell.font = bold_cell_font
                        else:
                            cell.fill = safety_no_fill
                            
                    if c_idx == 7: # Test Type
                        cell.alignment = align_center
                        tt_str = str(val).strip().upper()
                        if tt_str == 'NORMAL':
                            cell.fill = type_normal_fill
                        elif tt_str == 'DC':
                            cell.fill = type_dc_fill
                        elif tt_str == 'BOUNDARY':
                            cell.fill = type_boundary_fill
                m_row += 1

    # ----------------------------------------------------
    # SHEET 3: Requirements Catalog
    # ----------------------------------------------------
    ws_reqs = wb.create_sheet(title="Requirements Catalog")
    ws_reqs.views.sheetView[0].showGridLines = True
    ws_reqs.freeze_panes = 'A2'
    
    req_headers = [
        "Req #", "Requirement ID", "Requirement Text", "Notes",
        "Req Type", "Safety Impact", "Safety Rationale", "Linked Test Cases Count"
    ]
    for c_idx, h in enumerate(req_headers, start=1):
        cell = ws_reqs.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    r_row = 2
    for r in file_reqs:
        fill = alt_row_fill if r_row % 2 == 0 else white_fill
        row_vals = [
            r['req_num'],
            r['req_id'],
            r['req_text'],
            r['notes'] if r['notes'] else "None",
            r['req_type'],
            r['safety_impact'],
            r['safety_rationale'],
            len(r['test_cases'])
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_reqs.cell(row=r_row, column=c_idx, value=val)
            cell.font = cell_font
            cell.fill = fill
            cell.alignment = align_left
            cell.border = thin_border
            
            if c_idx in (1, 8): # Req #, count
                cell.alignment = align_center
            if c_idx == 6: # Safety Impact
                cell.alignment = align_center
                if str(val).strip().lower() == 'yes':
                    cell.fill = safety_yes_fill
                    cell.font = bold_cell_font
                else:
                    cell.fill = safety_no_fill
        r_row += 1

    # ----------------------------------------------------
    # SHEET 4: Test Cases Catalog
    # ----------------------------------------------------
    ws_tcs = wb.create_sheet(title="Test Cases Catalog")
    ws_tcs.views.sheetView[0].showGridLines = True
    ws_tcs.freeze_panes = 'A2'
    
    tc_headers = [
        "Requirement ID", "Test Case ID", "Test Type", "Description",
        "Initial Condition", "Test Inputs", "Expected Result", "Pass Criteria",
        "Related Requirements", "Test Procedure Notes"
    ]
    for c_idx, h in enumerate(tc_headers, start=1):
        cell = ws_tcs.cell(row=1, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = slate_fill
        cell.alignment = align_center
        cell.border = thin_border
        
    t_row = 2
    for r in file_reqs:
        for tc in r['test_cases']:
            fill = alt_row_fill if t_row % 2 == 0 else white_fill
            rel_reqs = tc.get('related_requirements', 'None')
            if isinstance(rel_reqs, list):
                rel_reqs = ", ".join(rel_reqs) if rel_reqs else "None"
                
            row_vals = [
                tc.get('requirement_id', r['req_id']),
                tc.get('test_case_id', 'N/A'),
                tc.get('test_type', 'N/A'),
                tc.get('description', ''),
                tc.get('initial_condition', ''),
                tc.get('test_inputs', ''),
                tc.get('expected_result', ''),
                tc.get('pass_criteria', ''),
                rel_reqs,
                tc.get('test_procedure_notes', '')
            ]
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_tcs.cell(row=t_row, column=c_idx, value=val)
                cell.font = cell_font
                cell.fill = fill
                cell.alignment = align_left
                cell.border = thin_border
                
                if c_idx == 3: # Test Type
                    cell.alignment = align_center
                    tt_str = str(val).strip().upper()
                    if tt_str == 'NORMAL':
                        cell.fill = type_normal_fill
                    elif tt_str == 'DC':
                        cell.fill = type_dc_fill
                    elif tt_str == 'BOUNDARY':
                        cell.fill = type_boundary_fill
            t_row += 1

    # Auto-adjust column widths
    for ws in [ws_dash, ws_matrix, ws_reqs, ws_tcs]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                lines = val_str.split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
            if max_len < 10:
                ws.column_dimensions[col_letter].width = 14
            elif max_len > 60:
                ws.column_dimensions[col_letter].width = 50
            else:
                ws.column_dimensions[col_letter].width = max_len + 4
                
    ws_matrix.column_dimensions['B'].width = 18 # Req ID
    ws_matrix.column_dimensions['C'].width = 45 # Req Text
    ws_matrix.column_dimensions['F'].width = 28 # TC ID
    ws_matrix.column_dimensions['H'].width = 35 # Description
    ws_matrix.column_dimensions['J'].width = 35 # Inputs
    ws_matrix.column_dimensions['K'].width = 40 # Expected Result
    ws_matrix.column_dimensions['L'].width = 30 # Pass Criteria
    
    ws_reqs.column_dimensions['C'].width = 50 # Req Text
    ws_reqs.column_dimensions['D'].width = 35 # Notes
    ws_reqs.column_dimensions['G'].width = 35 # Rationale
    
    ws_tcs.column_dimensions['B'].width = 28 # TC ID
    ws_tcs.column_dimensions['D'].width = 40 # Description
    ws_tcs.column_dimensions['F'].width = 35 # Inputs
    ws_tcs.column_dimensions['G'].width = 40 # Expected Result

    ws_dash.column_dimensions['A'].width = 4
    ws_dash.column_dimensions['B'].width = 24
    ws_dash.column_dimensions['C'].width = 35
    ws_dash.column_dimensions['D'].width = 24
    ws_dash.column_dimensions['E'].width = 20
    ws_dash.column_dimensions['F'].width = 24

    wb.save(output_file_path)
    print(f"Generated Excel: {output_file_path}")

def main():
    extracted_dir = 'extracted_samples'
    for fname in sorted(os.listdir(extracted_dir)):
        if fname.endswith('.txt'):
            fpath = os.path.join(extracted_dir, fname)
            with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
            reqs = parse_file_content(fname, content)
            
            # Form base output name: sample_1.xlsx or sample 1.xlsx
            base_name = os.path.splitext(fname)[0] # e.g. "sample 1"
            out_excel_name = f"{base_name}.xlsx" # e.g. "sample 1.xlsx"
            out_path = os.path.join('c:/Users/pm89542/Desktop/parse', out_excel_name)
            
            build_excel_for_file(fname, reqs, out_path)

if __name__ == '__main__':
    main()
