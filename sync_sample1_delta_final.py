import openpyxl
import json
import os

def sync_delta_final():
    excel_final_path = r"c:\Users\pm89542\Desktop\parse\SW_Requirements_Sample_1_Test_Cases_DELTA_FINAL.xlsx"
    out_dir = r"c:\Users\pm89542\Desktop\parse\app\ui_outputs"
    cache_dir = r"c:\Users\pm89542\Desktop\parse\app\cache"
    os.makedirs(out_dir, exist_ok=True)
    os.makedirs(cache_dir, exist_ok=True)

    json_out_path = os.path.join(out_dir, "SW_Requirements_Sample_1_generated_testcases.json")
    json_cache_path = os.path.join(cache_dir, "cache_SW_Requirements_Sample_1.json")
    excel_out_path = os.path.join(out_dir, "SW_Requirements_Sample_1_Test_Cases.xlsx")

    wb = openpyxl.load_workbook(excel_final_path)
    ws = wb.active

    test_cases = []
    cached_map = {}

    for r in range(3, ws.max_row + 1):
        sl_no = ws.cell(r, 1).value
        tc_id = str(ws.cell(r, 2).value or "").strip()
        req_id = str(ws.cell(r, 3).value or "").strip()
        test_type = str(ws.cell(r, 4).value or "NORMAL").strip()
        desc = str(ws.cell(r, 5).value or "").strip()
        init = str(ws.cell(r, 6).value or "").strip()
        inp = str(ws.cell(r, 7).value or "").strip()
        exp = str(ws.cell(r, 8).value or "").strip()
        criteria = str(ws.cell(r, 9).value or "").strip()
        comments = str(ws.cell(r, 10).value or "").strip()
        notes = str(ws.cell(r, 11).value or "").strip()

        if tc_id or req_id:
            tc_dict = {
                "test_case_id": tc_id,
                "requirement_id": req_id,
                "test_type": test_type,
                "description": desc,
                "initial_condition": init,
                "test_inputs": inp,
                "expected_result": exp,
                "pass_criteria": criteria,
                "related_requirements": comments if comments != "None" else "",
                "test_procedure_notes": notes
            }
            test_cases.append(tc_dict)
            cached_map.setdefault(req_id, []).append(tc_dict)

    # Save to JSON outputs and cache
    with open(json_out_path, "w", encoding="utf-8") as f:
        json.dump(test_cases, f, indent=2)

    with open(json_cache_path, "w", encoding="utf-8") as f:
        json.dump(cached_map, f, indent=2)

    # Save Excel output
    wb.save(excel_out_path)

    print(f"Successfully synced DELTA_FINAL data! Total Test Cases: {len(test_cases)}")

if __name__ == "__main__":
    sync_delta_final()
