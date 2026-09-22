import docx

doc = docx.Document('c:/Users/pm89542/Desktop/parse/SW_Requirements_Sample_1.docx')

print(f"Total tables in doc: {len(doc.tables)}")

for t_idx, table in enumerate(doc.tables):
    print(f"\n--- TABLE {t_idx+1} ---")
    for r_idx, row in enumerate(table.rows):
        cells_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
        row_str = " | ".join(cells_text)
        print(f"Row {r_idx}: {row_str.encode('ascii', errors='replace').decode('ascii')[:150]}")
