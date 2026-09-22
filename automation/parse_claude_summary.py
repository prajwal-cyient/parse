import openpyxl

path = r"c:\Users\pm89542\Desktop\parse\Sample_1\AI_vs_Manual_TestCase_Comparison_Req1001.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

print("WORKBOOK SHEETS:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n=========================================================================")
    print(f"SHEET: {sheet_name} (Rows: {ws.max_row}, Cols: {ws.max_column})")
    print(f"=========================================================================")
    
    # Print non-empty header/summary rows
    for r in range(1, min(40, ws.max_row + 1)):
        row_vals = [str(ws.cell(row=r, column=c).value).strip() if ws.cell(row=r, column=c).value is not None else "" for c in range(1, min(10, ws.max_column + 1))]
        if any(v for v in row_vals if v != "None"):
            print(f"Row {r:2d}: {' | '.join(row_vals[:6])}")
