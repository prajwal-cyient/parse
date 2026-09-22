import os
import re
from docx import Document

DOC_PATH = r"c:\Users\pm89542\Downloads\test_parker\SW_Requirements_Sample_1.docx"

doc = Document(DOC_PATH)

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

# Collect all Requirement IDs in document (e.g. LRU-xxxx or LRUSWRS-xxxx)
req_pattern = re.compile(r"LRU[A-Z0-9_-]+", re.IGNORECASE)

all_req_ids = set()

# Inspect paragraphs
text_reqs = []
for idx, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    matches = req_pattern.findall(text)
    if matches:
        for m in matches:
            all_req_ids.add(m)
            text_reqs.append((m, text))

# Inspect tables
table_info = []
for t_idx, table in enumerate(doc.tables, 1):
    table_reqs = []
    headers = [c.text.strip().replace("\n", " ") for c in table.rows[0].cells] if len(table.rows) > 0 else []
    for r_idx, row in enumerate(table.rows[1:], 2):
        row_text = " | ".join([c.text.strip().replace("\n", " ") for c in row.cells])
        matches = req_pattern.findall(row_text)
        for m in matches:
            all_req_ids.add(m)
            table_reqs.append((m, row_text))
            
    table_info.append({
        "table_num": t_idx,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "headers": headers,
        "req_count": len(table_reqs),
        "sample_reqs": table_reqs[:3]
    })

print("\n=========================================================================")
print(f"TOTAL UNIQUE REQUIREMENT IDs FOUND IN DOCUMENT: {len(all_req_ids)}")
print("=========================================================================\n")

print("TABLE-BY-TABLE REQUIREMENT BREAKDOWN:")
for t in table_info:
    print(f"Table {t['table_num']}: {t['rows']} rows x {t['cols']} cols | Header: {t['headers'][:3]}")
    print(f"   -> Requirements inside Table {t['table_num']}: {t['req_count']}")
    if t['sample_reqs']:
        for m, sample in t['sample_reqs'][:2]:
            print(f"      * {m}: {sample[:100]}...")
    print()

print("=========================================================================")
print(f"Paragraph text requirement references count: {len(text_reqs)}")
print("=========================================================================")
