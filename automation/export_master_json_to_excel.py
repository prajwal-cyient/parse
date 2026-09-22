import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

JSON_PATH = r"c:\Users\pm89542\Desktop\parse\automation\ALL_TABLE_1001_EXPANDED_TESTCASES.json"
EXCEL_PATH = r"c:\Users\pm89542\Desktop\parse\automation\ALL_TABLE_1001_EXPANDED_TESTCASES_FINAL.xlsx"

with open(JSON_PATH, "r", encoding="utf-8") as f:
    master_data = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "All 545 Expanded Test Cases"

# Ensure grid lines are visible
ws.views.sheetView[0].showGridLines = True

# Title Header Banner
ws.merge_cells("A1:K1")
title_cell = ws["A1"]
title_cell.value = "DO-178C LEVEL B TABLE 1001 HARDWARE-EXPANDED VERIFICATION TEST CASES (545 TOTAL)"
title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
title_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Navy Blue
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 40

# Table Headers
headers = [
    "Step",
    "Test Case ID",
    "Requirement Trace",
    "Test Type",
    "Test Case Description",
    "Initial Condition(s)",
    "Test Inputs",
    "Expected Result(s)",
    "Pass / Fail Criteria",
    "Related Requirements",
    "Test Procedure Notes"
]

header_fill = PatternFill(start_color="2B579A", end_color="2B579A", fill_type="solid") # Darker Blue
header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC")
)

ws.row_dimensions[2].height = 28

for col_num, header_title in enumerate(headers, 1):
    cell = ws.cell(row=2, column=col_num)
    cell.value = header_title
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = thin_border

# Color fills for Test Types
fill_normal = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft Green
fill_dc = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft Orange
fill_boundary = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Soft Yellow
fill_robustness = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid") # Soft Red
fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

font_body = Font(name="Arial", size=10, color="000000")
font_bold = Font(name="Arial", size=10, bold=True, color="000000")

row_idx = 3
steps_breakdown = master_data.get("steps_breakdown", {})

for step_name, step_info in steps_breakdown.items():
    step_id_clean = step_name.replace("STEP_", "Step ").replace("_REQUIREMENT_ID_", " (").replace("_", " ") + ")"
    req_id = step_info.get("requirement_id", "")
    tcs = step_info.get("test_cases", [])
    
    for tc in tcs:
        ttype = str(tc.get("test_type", "NORMAL")).upper()
        
        ws.cell(row=row_idx, column=1, value=step_id_clean).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=2, value=tc.get("test_case_id", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=3, value=tc.get("requirement_id", req_id)).alignment = Alignment(vertical="top", wrap_text=True)
        
        type_cell = ws.cell(row=row_idx, column=4, value=ttype)
        type_cell.alignment = Alignment(horizontal="center", vertical="top")
        type_cell.font = font_bold
        if "DC" in ttype:
            type_cell.fill = fill_dc
        elif "BOUND" in ttype:
            type_cell.fill = fill_boundary
        elif "ROBUST" in ttype:
            type_cell.fill = fill_robustness
        else:
            type_cell.fill = fill_normal
            
        ws.cell(row=row_idx, column=5, value=tc.get("description", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=6, value=tc.get("initial_condition", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=7, value=tc.get("test_inputs", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=8, value=tc.get("expected_result", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=9, value=tc.get("pass_criteria", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=10, value=tc.get("related_requirements", "")).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_idx, column=11, value=tc.get("test_procedure_notes", "")).alignment = Alignment(vertical="top", wrap_text=True)
        
        # Apply borders and font to row
        is_even = (row_idx % 2 == 0)
        for c in range(1, 12):
            cell = ws.cell(row=row_idx, column=c)
            cell.border = thin_border
            if c != 4: # Keep test_type fill
                if is_even:
                    cell.fill = fill_zebra
            if c in [1, 2, 3]:
                cell.font = font_bold
            else:
                cell.font = font_body
                
        ws.row_dimensions[row_idx].height = 42
        row_idx += 1

# Set optimal column widths
col_widths = {
    1: 22, # Step
    2: 32, # Test Case ID
    3: 16, # Requirement Trace
    4: 16, # Test Type
    5: 45, # Description
    6: 40, # Initial Condition
    7: 45, # Test Inputs
    8: 45, # Expected Result
    9: 35, # Pass/Fail Criteria
    10: 22, # Related Reqs
    11: 30  # Procedure Notes
}

for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

wb.save(EXCEL_PATH)

print("====================================================")
print("SUCCESSFULLY EXPORTED FINAL MASTER EXCEL SUITE!")
print(f"Total Rows Written: {row_idx - 3}")
print(f"Excel File Saved To: {EXCEL_PATH}")
print("====================================================")
