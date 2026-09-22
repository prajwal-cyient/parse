import openpyxl
import json
import os

base_dir = r"c:\Users\pm89542\Desktop\parse\samples"
ext_dir = r"c:\Users\pm89542\Desktop\parse\sampless_extracted"

samples_info = [
    ("Sample 2", "sample_2", "Sample_2", "sample_2_generated_testcases.json", "Test_Cases_For_Sample_2_Final.xlsx"),
    ("Sample 3", "sample_3", "Sample_3", "sample_3_generated_testcases.json", "Test_Cases_For_Sample_3.xlsx"),
    ("Sample 4", "sample_4", "Sample_4", "sample_4_generated_testcases.json", "Test_Cases_For_Sample_4.xlsx")
]

print("=========================================================================================================")
print("SIDE-BY-SIDE EXAMPLES: HUMAN MANUAL TEST CASES vs AI GENERATED TEST CASES")
print("=========================================================================================================\n")

for s_name, s_key, s_folder, json_name, excel_name in samples_info:
    json_path = os.path.join(base_dir, s_key, json_name)
    excel_path = os.path.join(ext_dir, s_folder, excel_name)
    
    with open(json_path, "r", encoding="utf-8") as f:
        gen_data = json.load(f)
    gen_tcs = gen_data.get("test_cases", [])
    
    wb_human = openpyxl.load_workbook(excel_path, data_only=True)
    ws_human = wb_human["Test Cases"] if "Test Cases" in wb_human.sheetnames else wb_human.active
    
    print(f"===================== {s_name.upper()} SIDE-BY-SIDE MATCHES =====================")
    
    matched_examples = []
    for r in range(2, ws_human.max_row + 1):
        h_id = str(ws_human.cell(row=r, column=1).value or "")
        h_obj = str(ws_human.cell(row=r, column=2).value or ws_human.cell(row=r, column=5).value or "")
        h_req = str(ws_human.cell(row=r, column=3).value or ws_human.cell(row=r, column=1).value or "")
        h_type = str(ws_human.cell(row=r, column=4).value or "NORMAL")
        h_inputs = str(ws_human.cell(row=r, column=6).value or ws_human.cell(row=r, column=7).value or "")
        h_exp = str(ws_human.cell(row=r, column=7).value or ws_human.cell(row=r, column=8).value or "")
        
        for g_tc in gen_tcs:
            g_req = g_tc.get("requirement_id", "")
            if g_req.upper() in h_req.upper() or g_req.upper() in h_obj.upper():
                matched_examples.append((h_id, h_req, h_obj, h_inputs, h_exp, g_tc))
                break
        if len(matched_examples) >= 2:
            break
            
    for idx, (h_id, h_req, h_obj, h_inp, h_exp, g_tc) in enumerate(matched_examples, 1):
        r_title = g_tc.get('requirement_id', '')
        g_tc_id = g_tc.get('test_case_id', '')
        g_desc = g_tc.get('description', '').strip()
        g_inp = g_tc.get('test_inputs', '').strip()
        g_out = g_tc.get('expected_result', '').strip()
        g_pass = g_tc.get('pass_criteria', '').strip()
        
        print(f"--- Example {idx} [{s_name} - Req: {r_title}] ---")
        print(f"[HUMAN MANUAL TEST CASE ID: {h_id}]")
        print(f"   Requirement Trace : {h_req[:70]}")
        print(f"   Objective/Desc    : {h_obj.strip()[:110]}")
        print(f"   Test Inputs       : {h_inp.strip()[:110]}")
        print(f"   Expected Result   : {h_exp.strip()[:110]}\n")
        
        print(f"[OUR AI GENERATED TEST CASE ID: {g_tc_id}]")
        print(f"   Requirement Trace : {r_title}")
        print(f"   Objective/Desc    : {g_desc[:110]}")
        print(f"   Test Inputs       : {g_inp[:110]}")
        print(f"   Expected Result   : {g_out[:110]}")
        print(f"   Pass Criteria     : {g_pass[:110]}\n")
        print("-" * 85)
