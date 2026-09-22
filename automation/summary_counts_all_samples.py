import openpyxl
import json
import os

manifest = [
    ("Sample 1", [r"c:\Users\pm89542\Downloads\test_parker\Test_Cases_For_Sample_1_Part_1.xlsx", r"c:\Users\pm89542\Downloads\test_parker\Test_Cases_For_Sample_1_Part_2.xlsx"], r"c:\Users\pm89542\Desktop\parse\automation\ALL_TABLE_1001_EXPANDED_TESTCASES.json"),
    ("Sample 2", [r"c:\Users\pm89542\Desktop\parse\sampless_extracted\Sample_2\Test_Cases_For_Sample_2_Final.xlsx"], r"c:\Users\pm89542\Desktop\parse\samples\sample_2\sample_2_generated_testcases.json"),
    ("Sample 3", [r"c:\Users\pm89542\Desktop\parse\sampless_extracted\Sample_3\Test_Cases_For_Sample_3.xlsx"], r"c:\Users\pm89542\Desktop\parse\samples\sample_3\sample_3_generated_testcases.json"),
    ("Sample 4", [r"c:\Users\pm89542\Desktop\parse\sampless_extracted\Sample_4\Test_Cases_For_Sample_4.xlsx"], r"c:\Users\pm89542\Desktop\parse\samples\sample_4\sample_4_generated_testcases.json")
]

print("=========================================================================================================")
print("TOTAL COMPARISON SUMMARY: HUMAN EXCEL TEST CASES vs OUR AI GENERATED TEST CASES")
print("=========================================================================================================\n")

grand_human = 0
grand_ai = 0

print(f"{'Sample Name':<12} | {'Human Excel Files Used':<55} | {'Human Ref TCs':<15} | {'Our AI Gen TCs':<15}")
print("-" * 105)

for s_name, excel_list, json_path in manifest:
    h_cnt = 0
    file_names = []
    for ep in excel_list:
        if os.path.exists(ep):
            wb = openpyxl.load_workbook(ep, data_only=True)
            ws = wb['Test Cases'] if 'Test Cases' in wb.sheetnames else wb.active
            rows = ws.max_row - 1
            h_cnt += rows
            file_names.append(f"{os.path.basename(ep)} ({rows} rows)")
            
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    if 'steps_breakdown' in data:
        ai_cnt = sum(len(step['test_cases']) for step in data['steps_breakdown'].values())
    else:
        ai_cnt = len(data.get('test_cases', []))
        
    grand_human += h_cnt
    grand_ai += ai_cnt
    
    files_str = ", ".join(file_names)
    print(f"{s_name:<12} | {files_str:<55} | {h_cnt:<15} | {ai_cnt:<15}")

print("-" * 105)
print(f"{'GRAND TOTAL':<12} | {'ALL 4 SAMPLES COMBINED':<55} | {grand_human:<15} | {grand_ai:<15}")
print("=========================================================================================================")
