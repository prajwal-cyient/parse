import openpyxl

path = r"c:\Users\pm89542\Desktop\parse\Sample_1\AI_vs_Manual_TestCase_Comparison_Req1001.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

print("WORKBOOK SHEETS:", wb.sheetnames)

first_sheet = wb[wb.sheetnames[0]]
print(f"\n=========================================================================")
print(f"READING FIRST SHEET: {wb.sheetnames[0]}")
print(f"=========================================================================")

for r in range(1, first_sheet.max_row + 1):
    row_vals = [str(first_sheet.cell(row=r, column=c).value).strip() if first_sheet.cell(row=r, column=c).value is not None else "" for c in range(1, first_sheet.max_column + 1)]
    non_empty = [v for v in row_vals if v != "" and v != "None"]
    if non_empty:
        print(f"Row {r:3d}: {' | '.join(non_empty)}")
