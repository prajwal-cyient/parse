import openpyxl, json, re

wb = openpyxl.load_workbook('app/outputs/SW_Requirements_Sample_1_Test_Cases_PERFECTED_APPROVED.xlsx')
ws = wb.active
headers = [c for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
tcs = []
for row in ws.iter_rows(min_row=3, values_only=True):
    tcs.append(dict(zip(headers, row)))

print(f"Total test cases to check: {len(tcs)}")

for i, tc in enumerate(tcs):
    tc_id = tc['Test Case ID']
    desc = str(tc['Test Case Description'] or '')
    inp = str(tc['Test Inputs'] or '')
    er = str(tc['Expected Result(s)'] or '')
    pc = str(tc['Pass / Fail Criteria'] or '')
    ic = str(tc['Initial Condition(s)'] or '')
    pn = str(tc['Test Procedure Notes'] or '')

    # 1. Check for wrong booleans in numeric variables
    for field_name, field_val in [('Inputs', inp), ('ER', er), ('PC', pc)]:
        for var in ['Harmonize_Offset_Default', 'Harmonize_Offset_New', 'Harmonize_SFC_Default', 'Act_Disp_Raw', 'K_Harmonize_New', 'Profile_Id']:
            m = re.search(rf'{var}\s*=\s*(True|False)', field_val, re.IGNORECASE)
            if m:
                print(f"[{tc_id}] Datatype issue in {field_name}: '{m.group(0)}' in {field_val}")

    # 2. Check for double assignments
    if '= false =' in er.lower() or '= true =' in er.lower() or '= false =' in inp.lower() or '= true =' in inp.lower():
        print(f"[{tc_id}] Double assignment: ER='{er}' | INP='{inp}'")

    # 3. Check for Pass Criteria alignment with Expected Result
    if 'equals true' in pc.lower() and ' = true' not in er.lower() and 'is true' not in er.lower() and 'true' not in er.lower():
        print(f"[{tc_id}] PC says True but ER doesn't have True: ER='{er}' | PC='{pc}'")

    if 'equals false' in pc.lower() and ' = false' not in er.lower() and 'is false' not in er.lower() and 'false' not in er.lower():
        print(f"[{tc_id}] PC says False but ER doesn't have False: ER='{er}' | PC='{pc}'")
