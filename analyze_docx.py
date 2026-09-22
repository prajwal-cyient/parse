from docx import Document
import sys
import json

def analyze_docx(path):
    doc = Document(path)
    
    analysis = {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sample_text": [p.text for p in doc.paragraphs if p.text.strip()][:10]
    }
    
    print(json.dumps(analysis, indent=2))

if __name__ == "__main__":
    analyze_docx(sys.argv[1])
